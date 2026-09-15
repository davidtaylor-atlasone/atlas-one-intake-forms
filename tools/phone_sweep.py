#!/usr/bin/env python3
"""Atlas One phone sweep: 385-213-7177 -> 380-225-5217 in docx / pptx / xlsx.
Run by David in the plain Terminal (Claude Code terminals are not allowed to overwrite OneDrive files).
Backs up every changed file to _to_delete/superseded-2026-09-14/phone-sweep/ first.
Signature / footer lines get "380-225-5217 (380-CALL-A1S)"; plain sentences get "380-225-5217".
"""
import os, re, shutil, sys, glob

ROOT = glob.glob(os.path.expanduser("~/Library/CloudStorage/OneDrive-AtlasOneSolutions/2. A1 Official Docs/2. Atlas 1 Solutions Marketing"))
if not ROOT:
    sys.exit("Marketing folder not found")
ROOT = ROOT[0]
BACKUP = os.path.join(os.path.dirname(os.path.dirname(ROOT)), "_to_delete", "superseded-2026-09-14", "phone-sweep")
SKIP = ("_to_delete", "_gold_backup", "3. Cornerstone PEO", "~$")

OLD = re.compile(r"(\+?1[\s.\-]?)?\(?385\)?[\s.\-]?213[\s.\-]?7177")
VANITY = "380-225-5217 (380-CALL-A1S)"
PLAIN = "380-225-5217"
SIG_HINT = re.compile(r"(atlasonesolutions|@|founder|president|lehi|book time|\|)", re.I)

def new_number(text):
    # Footer / signature style lines (email, website, title or pipes nearby) get the vanity form.
    return VANITY if SIG_HINT.search(text) else PLAIN

def fix_text(text):
    if not OLD.search(text):
        return text, False
    # Do not double up if the vanity text already follows the old number.
    text = re.sub(r"385[\s.\-]?213[\s.\-]?7177\s*\(380-CALL-A1S\)", "385-213-7177", text)
    return OLD.sub(new_number(text), text), True

def fix_runs(paragraph):
    """Replace inside a docx/pptx paragraph. Tries run by run first (keeps formatting);
    if the number is split across runs, rewrites the paragraph text into the first run."""
    changed = False
    for r in paragraph.runs:
        t, c = fix_text(r.text)
        if c:
            r.text = t; changed = True
    if not changed and OLD.search(paragraph.text):
        full, _ = fix_text(paragraph.text)
        for i, r in enumerate(paragraph.runs):
            r.text = full if i == 0 else ""
        changed = True
    return changed

def docx_paragraphs(doc):
    for p in doc.paragraphs: yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs: yield p
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            for p in part.paragraphs: yield p
            for t in part.tables:
                for row in t.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs: yield p

def do_docx(path):
    import docx
    d = docx.Document(path); changed = 0
    for p in docx_paragraphs(d):
        if fix_runs(p): changed += 1
    if changed: d.save(path)
    return changed

def do_pptx(path):
    from pptx import Presentation
    prs = Presentation(path); changed = 0
    def shapes_of(container):
        for sh in container.shapes:
            if sh.shape_type == 6:  # group
                yield from shapes_of(sh)
            else:
                yield sh
    for slide in list(prs.slides) + [prs.slide_master] + list(prs.slide_layouts):
        for sh in shapes_of(slide):
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    if fix_runs(p): changed += 1
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            if fix_runs(p): changed += 1
    if changed: prs.save(path)
    return changed

def do_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path); changed = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    t, ch = fix_text(c.value)
                    if ch: c.value = t; changed += 1
    if changed: wb.save(path)
    return changed

HANDLERS = {".docx": do_docx, ".pptx": do_pptx, ".xlsx": do_xlsx}

def main():
    dry = "--dry-run" in sys.argv
    fixed, errors = [], []
    for dirpath, dirnames, filenames in os.walk(ROOT):
        dirnames[:] = [d for d in dirnames if not any(s in d for s in SKIP)]
        for name in filenames:
            ext = os.path.splitext(name)[1].lower()
            if ext not in HANDLERS or name.startswith("~$"):
                continue
            path = os.path.join(dirpath, name)
            rel = os.path.relpath(path, ROOT)
            backup_path = os.path.join(BACKUP, rel)
            try:
                if dry:
                    # Read-only pass: copy to a temp file and test there.
                    import tempfile
                    tmp = os.path.join(tempfile.gettempdir(), name)
                    shutil.copy2(path, tmp)
                    n = HANDLERS[ext](tmp)
                    os.remove(tmp)
                else:
                    os.makedirs(os.path.dirname(backup_path), exist_ok=True)
                    if not os.path.exists(backup_path):
                        shutil.copy2(path, backup_path)
                    n = HANDLERS[ext](path)
                    if not n:
                        os.remove(backup_path)
                if n:
                    fixed.append((rel, n))
                    print(f"{'WOULD FIX' if dry else 'FIXED'} ({n} places): {rel}")
            except Exception as e:
                errors.append((rel, str(e)[:120]))
    print(f"\n{'Would change' if dry else 'Changed'} {len(fixed)} files. Errors: {len(errors)}")
    for rel, err in errors:
        print(f"  could not open: {rel}  ({err})")
    if fixed and not dry:
        print(f"\nOriginals backed up to: {BACKUP}")
        print("PDFs next to these files still show the old number; re-export them from the fixed source.")

if __name__ == "__main__":
    main()
