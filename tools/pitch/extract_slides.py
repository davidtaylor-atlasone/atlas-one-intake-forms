#!/usr/bin/env python3
"""Extracts every slide's text from the 16-slide Prospect Pitch Deck into
slides.json (Run AT, 2026-09-15). build.py reads slides.json, never the
pptx directly, so re-run this whenever the deck changes and then re-run
build.py to regenerate index.html.

Run:  python3 extract_slides.py "/path/to/Atlas_One_Master_Kit"
"""
import io
import json
import os
import sys

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))


def main():
    mk = sys.argv[1] if len(sys.argv) > 1 else "."
    mkt = os.path.normpath(os.path.join(mk, "..", ".."))
    path = os.path.join(mkt, "A1_Sales", "A1_Pitch Decks Inv", "Atlas_One_Prospect_Pitch_Deck.pptx")

    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        title = None
        body = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if not text.strip():
                continue
            is_title = (shape == slide.shapes.title) if slide.shapes.title else False
            if is_title and title is None:
                title = text.strip()
            else:
                for line in text.split("\n"):
                    line = line.strip()
                    if line:
                        body.append(line)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"index": i, "title": title or "", "body": body, "notes": notes})

    out_path = os.path.join(HERE, "slides.json")
    with io.open(out_path, "w", encoding="utf-8") as f:
        json.dump(slides, f, indent=2)
    print(f"wrote {out_path}, {len(slides)} slides")


if __name__ == "__main__":
    main()
