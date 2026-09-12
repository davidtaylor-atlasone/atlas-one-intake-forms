#!/usr/bin/env python3
import pptx, copy, sys, json
from pptx.dml.color import RGBColor
from pptx.util import Pt, Emu
from PIL import ImageFont
SRC='deck_in.pptx'; OUT='deck_fixed.pptx'; CUT='deck_cut.pptx'
HORAS="/Users/davidtaylor/Library/CloudStorage/OneDrive-AtlasOneSolutions/2. A1 Official Docs/2. Atlas 1 Solutions Marketing/HR_Docs/Atlas_One_Master_Kit/A1_Final Brand/3. Fonts/Horas-Medium.ttf"
NAVY=RGBColor(0x23,0x30,0x4D); PERI=RGBColor(0x78,0x8D,0xE3); SOFT=RGBColor(0xDB,0xE4,0xED); OFF=RGBColor(0xFA,0xFA,0xF8); WHITE=RGBColor(0xFF,0xFF,0xFF)
GREEN={'059669','10B981'}; RED={'991B1B','EF4444'}; REDTINT={'FEF2F2','FECACA'}
E=9525.0
log={'recolored':0,'titles':[], 'white_text':0}
p=pptx.Presentation(SRC)
def rgb(c):
    try: return str(c.rgb)
    except Exception: return None
def inside(a,b):  # a inside b
    return a.left>=b.left-1 and a.top>=b.top-1 and a.left+a.width<=b.left+b.width+1 and a.top+a.height<=b.top+b.height+1
for si,s in enumerate(p.slides,1):
    navy_boxes=[]
    for sh in s.shapes:
        # fills
        try:
            if sh.fill.type==1:
                c=rgb(sh.fill.fore_color)
                thin=min(sh.width,sh.height)<=Emu(80000)
                if c in GREEN:
                    if thin: sh.fill.fore_color.rgb=PERI
                    else: sh.fill.fore_color.rgb=NAVY; navy_boxes.append(sh)
                    log['recolored']+=1
                elif c in RED:
                    sh.fill.fore_color.rgb=PERI if thin else SOFT; log['recolored']+=1
                elif c in REDTINT:
                    sh.fill.fore_color.rgb=OFF; log['recolored']+=1
        except Exception: pass
        # lines
        try:
            if sh.line.fill.type==1:
                c=rgb(sh.line.color)
                if c in GREEN: sh.line.color.rgb=NAVY; log['recolored']+=1
                elif c in RED or c in REDTINT: sh.line.color.rgb=SOFT; log['recolored']+=1
        except Exception: pass
        # fonts
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    try:
                        c=rgb(r.font.color) if r.font.color and r.font.color.type is not None else None
                    except Exception: c=None
                    if c in GREEN or c in RED: r.font.color.rgb=NAVY; log['recolored']+=1
    # text sitting on a box that is now navy -> white
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() and any(inside(sh,b) and sh is not b for b in navy_boxes):
            for para in sh.text_frame.paragraphs:
                for r in para.runs: r.font.color.rgb=WHITE; log['white_text']+=1
    # title overflow: Horas runs
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        runs=[r for para in sh.text_frame.paragraphs for r in para.runs]
        if not runs or not any((r.font.name or '')=='Horas' for r in runs): continue
        size=max((r.font.size.pt for r in runs if r.font.size), default=None)
        if not size: continue
        text=sh.text_frame.text
        boxw=sh.width/E-14.4; boxh=sh.height/E-7.2
        def fits(pt):
            f=ImageFont.truetype(HORAS,int(round(pt*96/72)))
            lines=0
            for para_text in text.split('\n'):
                words=para_text.split(' '); cur=''
                for w in words:
                    t=(cur+' '+w).strip()
                    if f.getlength(t)<=boxw: cur=t
                    else: lines+=1; cur=w
                lines+=1
            return (lines==1) or (lines*pt*96/72*1.2<=boxh+2), lines
        ok,ln=fits(size); new=size
        while not ok and new>size*0.6:
            new-=1; ok,ln=fits(new)
        if new!=size:
            for r in runs:
                if r.font.size: r.font.size=Pt(new)
            log['titles'].append((si,text[:50],size,new,ln))
p.save(OUT)
# ---- room cut ----
KEEP=[1,2,3,4,5,7,8,11,15,19,22,24,28,31,32,33,34,40,41,43]
q=pptx.Presentation(OUT)
sldIdLst=q.slides._sldIdLst
ids=list(sldIdLst)
for idx,sld in enumerate(ids,1):
    if idx not in KEEP:
        rId=sld.rId; q.part.drop_rel(rId); sldIdLst.remove(sld)
q.save(CUT)
print(json.dumps(log,indent=1)); print('cut slides',len(q.slides))
