#!/usr/bin/env python3
"""Run V job 1: slide 25 of the v9 LIGHT partner deck to the confirmed tiers. Three columns kept, Concierge as a one line note."""
import copy, shutil, os, sys
import pptx
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
DECK=sys.argv[1]; BACKUP=sys.argv[2]
os.makedirs(os.path.dirname(BACKUP),exist_ok=True)
if not os.path.exists(BACKUP): shutil.copy2(DECK,BACKUP)
p=pptx.Presentation(DECK); s=p.slides[24]
sh={x.name:x for x in s.shapes}
def settext(name,new):
    r=sh[name].text_frame.paragraphs[0].runs; r[0].text=new
    for extra in r[1:]: extra.text=''
    print('text',name,'->',new)
settext('Text 22','$399'); settext('Text 37','$999')
settext('Text 9','Bundled with software, setup waived')
settext('Text 23','Per company / per month, setup $495, often waived')
settext('Text 38','Per company / per month, setup $995, waived on annual')
settext('Text 2','Four Tiers + À La Carte Add-Ons.')
# sub lines are long now: 9pt -> 8pt so they stay on one line
for n in ('Text 9','Text 23','Text 38'):
    for r in sh[n].text_frame.paragraphs[0].runs: r.font.size=Pt(8)
# tighten the three cards by 0.2in to make room for the Concierge line
SHRINK=182880
for n in ('Shape 6','Shape 20','Shape 35'): sh[n].height=Emu(sh[n].height-SHRINK)
rows=[('Text 13','Text 14','Text 15','Text 16','Text 17'),
      ('Text 27','Text 28','Text 29','Text 30','Text 31','Text 32'),
      ('Text 42','Text 43','Text 44','Text 45','Text 46','Text 47')]
for col in rows:
    for i,n in enumerate(col):
        sh[n].top=Emu(3977640+i*228600); sh[n].height=Emu(219456)
# Concierge note: clone the lede text box (Text 3) for its styling, then restyle
src=sh['Text 3']._element; el=copy.deepcopy(src); src.getparent().append(el)
note=[x for x in s.shapes if x._element is el][0]
note.name='Text 55 Concierge note'
note.left=Emu(457200); note.top=Emu(5394960); note.width=Emu(11247120); note.height=Emu(228600)
para=note.text_frame.paragraphs[0]; para.alignment=PP_ALIGN.CENTER
r=para.runs[0]; r.text='Concierge, your fractional COO, $1,900 a month.'
r.font.size=Pt(11); r.font.italic=False; r.font.bold=True; r.font.color.rgb=RGBColor(0x23,0x30,0x4D)
# unique shape id
ids=[int(x._element.nvSpPr.cNvPr.get('id')) for x in s.shapes if x._element is not el]
el.nvSpPr.cNvPr.set('id',str(max(ids)+1))
p.save(DECK); print('saved',DECK)
