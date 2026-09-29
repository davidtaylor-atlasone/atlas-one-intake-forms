#!/usr/bin/env python3
"""Run AG3: load master-book-v5 into GoHighLevel (contacts + notes only).
Reads GHL_PIT from .env. Writes real data in every mode except `plan` -- read
BRIEF-AGENT-2026-09-27.md (AG3 section) before touching this.

Importable: job_ag3_sf_load.py reuses api(), throttle(), normalize_phone(), domain_of(),
search_by_email(), search_by_phone(), get_contact(), get_notes(), has_note_with_title(),
post_note_idempotent(), mask_pii(), blank(), split_name(), load_csv_log(), append_log(),
update_contact_fields(), add_tags(), create_contact(), set_dnd_email_only().
"""
import csv
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request

import openpyxl

REPO = "/Users/davidtaylor/Projects/atlas-one-intake-forms"
MK = "/Users/davidtaylor/Library/CloudStorage/OneDrive-AtlasOneSolutions/2. A1 Official Docs/2. Atlas 1 Solutions Marketing/HR_Docs/Atlas_One_Master_Kit"
BOOK_PATH = f"{MK}/_INTERNAL (do not share)/Book of Business/master-book-v5-2026-09-28.xlsx"
LOG_PATH = f"{MK}/_INTERNAL (do not share)/Book of Business/ag2-load-log-2026-09-27.csv"

LOCATION_ID = "AzTPxnK2vSUj19jYoDmR"
DAVID_USER_ID = "vTV2wRivyR9f9XWNook3"
API_BASE = "https://services.leadconnectorhq.com"
API_VERSION = "2021-07-28"

CF_EMPLOYEE_COUNT = "jhys6ULyKSEL8pxXTdYt"
CF_CURRENT_PAYROLL_PROVIDER = "2Wg0Mpj4wO8GaJjZaqdo"

TAG_IMPORT = "book-import-2026-09"
NOTE_TITLE = "Book of business import 2026-09"
STATUS_TAG = {
    "Cornerstone Client": "book-cs-client",
    "G&A Client": "book-ga-client",
    "Former Teamworks client": "book-former-teamworks",
    "Former Client": "book-former-client",
    "Prospect open": "book-prospect",
    "Lost": "book-lost",
    "Atlas One Client": "book-atlas-client",
    "Broker or Partner": "book-partner",
    "Vendor": "book-vendor",
}
CLIENT_CURRENT_STATUSES = {"Cornerstone Client", "G&A Client", "Atlas One Client"}

# Rule 7 (AG2): referral partner domains. peoplepayglobal.com is named explicitly in the
# brief; the rest are every Broker/Partner/Vendor website or email domain found in the Book
# tab and the Atlas_One_Vendor_Partner_Directory.xlsx. Static set built once, not pulled
# live from GHL -- logged as an Assumption in every report that touches it.
PARTNER_DOMAINS = {
    "peoplepayglobal.com", "microsoft.com", "one-cfo.com", "simpliverified.com",
    "smartbenefits.co", "deel.com", "bigredjelly.com", "quickbooks.intuit.com",
    "intuit.com", "cornerstonepeo.com", "gnapartners.com", "verohcm.com",
    "rowrowrow.io", "rhinoinsurance.com", "thecfpartners.com", "excelhealthplans.com",
    "amplifiedresourcegroup.com", "redirecthealth.com", "newjourneconsulting.com",
    "hoffmanandcompany.com", "experiencedpayrollpros.com", "apextax.net",
    "goreboot.com",
}

# AG3 fix pass: rows from the AG2 dryrun10 test (identified by company name, not row
# number -- v5 renumbered many rows versus v4). --fix reruns exactly these 10 to fill
# companyName (and, for 5G Hearing, attach the contact the Book itself has none for).
FIX_COMPANIES = {
    "5G Hearing LLC", "AAMCOR Inc", "CIRQUE LODGE INC", "Ibex Plumbing",
    "1 2 3 STITCH", "1st Rate Mortgage", "24 Hour Express", "2Brothers Furniture",
    "5B USA LLC", "A Krete Inc",
}
FIX_CONTACT_OVERRIDES = {
    # Brief 2026-09-28 update: 5G Hearing has no contact in the Book; its Cornerstone
    # note names the owners. Use Paul Campoamor as the contact in the --fix pass only.
    "5G Hearing LLC": {"name": "Paul Campoamor", "email": "paul@amhchearing.com", "phone": "352-406-1985"},
}

FEIN_KEYWORD_RE = re.compile(r"\b(fein|ein|tax id)\b", re.I)
FEIN_NUM_RE = re.compile(r"\b(\d{2})-?(\d{7})\b")
SSN_RE = re.compile(r"\b\d{3}-\d{2}-\d{4}\b")


def load_env():
    env = {}
    with open(os.path.join(REPO, ".env")) as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip()
    return env


TOKEN = load_env()["GHL_PIT"]

HEADERS = {
    "Authorization": f"Bearer {TOKEN}",
    "Version": API_VERSION,
    "Content-Type": "application/json",
    "Accept": "application/json",
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36",
}


def api(method, path, params=None, body=None, retries=5):
    url = f"{API_BASE}{path}"
    if params:
        url += "?" + urllib.parse.urlencode(params)
    data = json.dumps(body).encode() if body is not None else None
    for attempt in range(retries):
        req = urllib.request.Request(url, data=data, headers=HEADERS, method=method)
        try:
            with urllib.request.urlopen(req) as resp:
                raw = resp.read().decode()
                return resp.status, (json.loads(raw) if raw else {})
        except urllib.error.HTTPError as e:
            raw = e.read().decode()
            try:
                parsed = json.loads(raw)
            except Exception:
                parsed = {"raw": raw}
            if e.code == 429 and attempt < retries - 1:
                time.sleep(2 ** attempt)
                continue
            return e.code, parsed
        except urllib.error.URLError:
            if attempt < retries - 1:
                time.sleep(2 ** attempt)
                continue
            raise
    raise RuntimeError("unreachable")


def throttle():
    time.sleep(0.22)  # ~4.5 calls/sec


def normalize_phone(raw):
    if not raw:
        return None
    digits = re.sub(r"[^\d]", "", str(raw))
    if not digits:
        return None
    if len(digits) == 10:
        return "+1" + digits
    if len(digits) == 11 and digits.startswith("1"):
        return "+" + digits
    if str(raw).strip().startswith("+"):
        return str(raw).strip()
    return "+" + digits


def domain_of(email_or_url):
    if not email_or_url:
        return None
    s = str(email_or_url).lower()
    m = re.search(r"@([\w.-]+)", s)
    if m:
        return m.group(1)
    m = re.search(r"([\w-]+\.[a-z]{2,})(?:/|$)", s.replace("http://", "").replace("https://", "").replace("www.", ""))
    if m:
        return m.group(1)
    return None


def parse_address(addr):
    if not addr:
        return None, None, None, None
    parts = [p.strip() for p in str(addr).split(",")]
    if len(parts) >= 4:
        return parts[0], parts[1], parts[2], parts[3]
    if len(parts) == 3:
        return parts[0], parts[1], None, parts[2]
    return addr, None, None, None


def mask_pii(text):
    """Mask FEIN/EIN/Tax Id shapes (only near that keyword) and SSN shapes (always)."""
    if not text:
        return text
    out_lines = []
    for line in text.split("\n"):
        if FEIN_KEYWORD_RE.search(line):
            def repl(m):
                digits = m.group(1) + m.group(2)
                return f"XX-XXX{digits[-4:]}"
            line = FEIN_NUM_RE.sub(repl, line)
        line = SSN_RE.sub(lambda m: f"XXX-XX-{re.sub(r'[^0-9]', '', m.group(0))[-4:]}", line)
        out_lines.append(line)
    return "\n".join(out_lines)


def search_by_email(email):
    if not email:
        return None
    status, data = api("GET", "/contacts/search/duplicate", params={"locationId": LOCATION_ID, "email": email})
    throttle()
    if status == 200:
        return data.get("contact")
    return None


def search_by_phone(phone):
    if not phone:
        return None
    status, data = api("GET", "/contacts/search/duplicate", params={"locationId": LOCATION_ID, "number": phone})
    throttle()
    if status == 200:
        return data.get("contact")
    return None


def get_contact(contact_id):
    if not contact_id:
        return None
    status, data = api("GET", f"/contacts/{contact_id}")
    throttle()
    if status == 200:
        return data.get("contact")
    return None


def get_notes(contact_id):
    if not contact_id:
        return []
    status, data = api("GET", f"/contacts/{contact_id}/notes")
    throttle()
    if status == 200:
        return data.get("notes", []) or []
    return []


def has_note_with_title(contact_id, title):
    return any((n.get("title") or "") == title for n in get_notes(contact_id))


def post_note_idempotent(contact_id, title, body, dry_run=False):
    """Returns (status_str, detail). status_str one of ok-added / ok-skip-dup / failed."""
    if has_note_with_title(contact_id, title):
        return "ok-skip-dup", "note with this title already exists, skipped"
    if dry_run:
        return "ok-would-add", f"would add note len={len(body)}"
    status_code, resp = api("POST", f"/contacts/{contact_id}/notes", body={"body": body, "title": title})
    throttle()
    if status_code not in (200, 201):
        return "failed", f"note failed {status_code}: {json.dumps(resp)[:300]}"
    return "ok-added", "added"


def update_contact_fields(contact_id, update_body, dry_run=False):
    if not update_body:
        return "ok-nothing", "no blank fields to fill"
    if dry_run:
        return "ok-would-update", f"would update {list(update_body.keys())}"
    status_code, resp = api("PUT", f"/contacts/{contact_id}", body=update_body)
    throttle()
    if status_code not in (200, 201):
        return "failed", f"update fields failed {status_code}: {json.dumps(resp)[:300]}"
    return "ok-updated", "updated"


def add_tags(contact_id, new_tags, dry_run=False):
    if not new_tags:
        return "ok-nothing", "no new tags"
    if dry_run:
        return "ok-would-tag", f"would add tags {new_tags}"
    status_code, resp = api("POST", f"/contacts/{contact_id}/tags", body={"tags": new_tags})
    throttle()
    if status_code not in (200, 201):
        return "failed", f"tags failed {status_code}: {json.dumps(resp)[:300]}"
    return "ok-tagged", "tagged"


def create_contact(create_body, dry_run=False):
    if dry_run:
        return "ok-would-create", None, f"would create {json.dumps(create_body)[:300]}"
    status_code, resp = api("POST", "/contacts/", body=create_body)
    throttle()
    if status_code not in (200, 201):
        return "failed", None, f"create failed {status_code}: {json.dumps(resp)[:300]}"
    return "ok-created", resp["contact"]["id"], "created"


def set_dnd_email_only(contact_id, dry_run=False):
    """David's approved exception (AG3, Salesforce load): DND on the EMAIL channel only,
    for contacts who unsubscribed from email in Salesforce. Never touches any other channel."""
    body = {"dndSettings": {"Email": {"status": "active", "message": "", "code": "OptOut"}}}
    if dry_run:
        return "ok-would-dnd-email", "would set dndSettings.Email.status = active"
    status_code, resp = api("PUT", f"/contacts/{contact_id}", body=body)
    throttle()
    if status_code not in (200, 201):
        return "failed", f"dnd-email failed {status_code}: {json.dumps(resp)[:300]}"
    return "ok-dnd-email", "dnd email set"


def blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def build_notes_index(wb):
    ws = wb["Notes"]
    idx = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        company, date, source, text = r
        if not company:
            continue
        idx.setdefault(company, []).append((date, source, text))
    for k in idx:
        idx[k].sort(key=lambda t: (t[0] or ""), reverse=True)
    return idx


def build_contacts_index(wb):
    ws = wb["Contacts"]
    idx = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        company, name, email, phone, source, note = r
        if not company:
            continue
        idx.setdefault(company, []).append({"name": name, "email": email, "phone": phone, "source": source, "note": note})
    return idx


def build_company_note(row_dict, notes_for_company, extra_note=None):
    lines = [
        f"Status: {row_dict['Status'] or ''}",
        f"PEO: {row_dict['PEO'] or ''}",
        f"Cornerstone entity code: {row_dict['Cornerstone entity code'] or ''}",
        f"CSR: {row_dict['CSR'] or ''}",
        f"Underwriter: {row_dict['Underwriter'] or ''}",
        f"Contract type: {row_dict['Contract type'] or ''}",
        f"Referral partner: {row_dict['Referral partner'] or ''}",
        f"Effective date: {row_dict['Effective date'] or ''}",
        f"Services: {row_dict['Services'] or ''}",
        f"Placement note: {row_dict['Placement note'] or ''}",
        f"Folder path(s): {row_dict['Folder path(s)'] or ''}",
    ]
    if extra_note:
        lines.append(extra_note)
    lines.append("")
    lines.append("Linked notes (newest first):")
    for date, source, text in notes_for_company:
        lines.append(f"[{date}] ({source}) {text}")
    body = "\n".join(lines)
    body = mask_pii(body)
    if len(body) > 60000:
        body = body[:59900] + "\n...(truncated at 60,000 characters)"
    return body


def pick_main_contact(row_dict, contacts_for_company):
    name = row_dict["Main contact name"]
    email = row_dict["Main contact email"]
    phone = row_dict["Main contact phone"]
    extra_note = None
    d = domain_of(email) if email else None
    if email and d in PARTNER_DOMAINS:
        skipped = f"Referral partner contact: {name}, {email}"
        alt = None
        for c in contacts_for_company:
            cd = domain_of(c.get("email"))
            if c.get("email") and cd not in PARTNER_DOMAINS:
                alt = c
                break
        if alt:
            name, email, phone = alt.get("name"), alt.get("email"), alt.get("phone")
        else:
            name, email = None, None
            phone = row_dict["Phone"]
        extra_note = skipped
    return name, email, phone, extra_note


def split_name(name):
    if not name:
        return None, None
    parts = str(name).strip().split()
    if len(parts) == 1:
        return parts[0], None
    return parts[0], " ".join(parts[1:])


def load_csv_log(log_path=LOG_PATH):
    done = set()
    if os.path.exists(log_path):
        with open(log_path, newline="") as f:
            for row in csv.DictReader(f):
                if row.get("result", "").startswith("ok"):
                    done.add(row["row"])
    return done


def append_log(row_key, company, action, contact_id, result, log_path=LOG_PATH):
    is_new = not os.path.exists(log_path)
    with open(log_path, "a", newline="") as f:
        w = csv.writer(f)
        if is_new:
            w.writerow(["row", "company", "action", "contact_id", "result"])
        w.writerow([row_key, company, action, contact_id or "", result])



# ---- Cowork 2026-09-28: shared phone guard -------------------------------------------------
# A phone number that appears on two or more different companies in the Book (for example
# PeoplePayGlobal's shared USA line +1 212 424 6015 on six PPG referred clients) must never be
# used to find an existing GHL contact, or unrelated companies get merged into one record.
PPG_SHARED_LINES = {"2124246015"}
_SHARED = None
def _digits10(p):
    d = re.sub(r"\D", "", str(p or ""))
    return d[1:] if len(d) == 11 and d.startswith("1") else d
def _norm_co(n):
    n = re.sub(r"[^a-z0-9 ]", " ", str(n or "").lower().replace("&", " and "))
    n = re.sub(r"\b(llc|inc|corp|co|company|pllc|pc|ltd|lp|llp|dba|the|group)\b", " ", n)
    return " ".join(n.split())
def shared_phones():
    global _SHARED
    if _SHARED is None:
        import collections
        seen = collections.defaultdict(set)
        wb = openpyxl.load_workbook(BOOK_PATH, read_only=True, data_only=True)
        rows = list(wb["Book"].iter_rows(values_only=True)); hdr = rows[0]
        ix = {h: i for i, h in enumerate(hdr)}
        for r in rows[1:]:
            for k in ("Phone", "Main contact phone"):
                d = _digits10(r[ix[k]]) if k in ix else ""
                if len(d) == 10 and r[ix["Company"]]:
                    seen[d].add(_norm_co(r[ix["Company"]]))
        _SHARED = {d for d, cos in seen.items() if len(cos) > 1} | PPG_SHARED_LINES
    return _SHARED


def build_plan(row_dict, notes_idx, contacts_idx, fix_pass=False):
    """Read-only: computes what would happen. Never writes. Used by plan mode and as the
    first half of every real-write pass so the decision logic lives in exactly one place."""
    company = row_dict["Company"]
    contacts_for_company = contacts_idx.get(company, [])

    override = FIX_CONTACT_OVERRIDES.get(company) if fix_pass else None
    if override:
        name, email, phone, extra_note = override["name"], override["email"], override["phone"], None
    else:
        name, email, phone, extra_note = pick_main_contact(row_dict, contacts_for_company)

    phone_norm = normalize_phone(phone) or normalize_phone(row_dict["Phone"])
    first, last = split_name(name)
    no_contact_name = not first and not last

    existing = None
    match_method = None
    if email:
        existing = search_by_email(email)
        if existing:
            match_method = "email"
    if not existing and phone_norm and _digits10(phone_norm) not in shared_phones():
        existing = search_by_phone(phone_norm)
        if existing:
            match_method = "phone"
            # Cowork 2026-09-28: a phone match to a contact that already carries a DIFFERENT
            # company name is a shared number (referral partner line, related entity), not the same
            # company. Treat it as no match and create a new contact instead of merging.
            ex_co = (existing.get("companyName") or "").strip()
            if ex_co and _norm_co(ex_co) != _norm_co(company):
                existing, match_method = None, None
    if not existing and row_dict.get("GHL Contact ID"):
        existing = get_contact(row_dict["GHL Contact ID"])
        if existing:
            match_method = "GHL Contact ID column"

    addr1, city, state, postal = parse_address(row_dict["Address"])
    if not state and row_dict.get("State"):
        state = row_dict["State"]

    status = row_dict["Status"]
    status_tag = STATUS_TAG.get(status)
    tags = [TAG_IMPORT] + ([status_tag] if status_tag else [])
    if status in CLIENT_CURRENT_STATUSES:
        tags.append("client-current")
    if no_contact_name:
        tags.append("no-contact-name")
    if not blank(row_dict.get("Needs your review")):
        tags.append("status-review")

    custom_fields = []
    employees = row_dict.get("Employees")
    if employees not in (None, ""):
        custom_fields.append({"id": CF_EMPLOYEE_COUNT, "fieldValue": employees})

    note_body = build_company_note(row_dict, notes_idx.get(company, []), extra_note)

    plan = {
        "company": company, "existing": existing, "match_method": match_method,
        "first": first, "last": last, "email": email, "phone_norm": phone_norm,
        "addr1": addr1, "city": city, "state": state, "postal": postal,
        "website": row_dict.get("Website"), "tags": tags, "custom_fields": custom_fields,
        "note_body": note_body, "company_name": company,
    }
    return plan


def execute_plan(plan, dry_run=False):
    """Executes (or, if dry_run, simulates) exactly what build_plan decided."""
    existing = plan["existing"]
    if existing:
        contact_id = existing["id"]
        update_body = {}
        if blank(existing.get("firstName")) and plan["first"]:
            update_body["firstName"] = plan["first"]
        if blank(existing.get("lastName")) and plan["last"]:
            update_body["lastName"] = plan["last"]
        if blank(existing.get("email")) and plan["email"]:
            update_body["email"] = plan["email"]
        if blank(existing.get("phone")) and plan["phone_norm"]:
            update_body["phone"] = plan["phone_norm"]
        if blank(existing.get("address1")) and plan["addr1"]:
            update_body["address1"] = plan["addr1"]
        if blank(existing.get("city")) and plan["city"]:
            update_body["city"] = plan["city"]
        if blank(existing.get("state")) and plan["state"]:
            update_body["state"] = plan["state"]
        if blank(existing.get("postalCode")) and plan["postal"]:
            update_body["postalCode"] = plan["postal"]
        if blank(existing.get("website")) and plan["website"]:
            update_body["website"] = plan["website"]
        if blank(existing.get("companyName")) and plan["company_name"]:
            update_body["companyName"] = plan["company_name"]
        existing_cf = {cf.get("id"): cf.get("value") for cf in existing.get("customFields", [])}
        fill_cf = [cf for cf in plan["custom_fields"] if blank(existing_cf.get(cf["id"]))]
        if fill_cf:
            update_body["customFields"] = fill_cf

        status1, detail1 = update_contact_fields(contact_id, update_body, dry_run=dry_run)
        if status1 == "failed":
            return "failed", contact_id, detail1

        new_tags = [t for t in plan["tags"] if t not in (existing.get("tags") or [])]
        status2, detail2 = add_tags(contact_id, new_tags, dry_run=dry_run)
        if status2 == "failed":
            return "failed", contact_id, detail2

        status3, detail3 = post_note_idempotent(contact_id, NOTE_TITLE, plan["note_body"], dry_run=dry_run)
        if status3 == "failed":
            return "failed", contact_id, detail3

        return "ok-updated" if not dry_run else "ok-would-update", contact_id, f"{detail1}; {detail2}; {detail3}"
    else:
        create_body = {"locationId": LOCATION_ID, "assignedTo": DAVID_USER_ID, "tags": plan["tags"],
                        "source": "Book of business import 2026-09", "companyName": plan["company_name"]}
        if plan["first"]:
            create_body["firstName"] = plan["first"]
        if plan["last"]:
            create_body["lastName"] = plan["last"]
        if plan["email"]:
            create_body["email"] = plan["email"]
        if plan["phone_norm"]:
            create_body["phone"] = plan["phone_norm"]
        if plan["addr1"]:
            create_body["address1"] = plan["addr1"]
        if plan["city"]:
            create_body["city"] = plan["city"]
        if plan["state"]:
            create_body["state"] = plan["state"]
        if plan["postal"]:
            create_body["postalCode"] = plan["postal"]
        if plan["website"]:
            create_body["website"] = plan["website"]
        if plan["custom_fields"]:
            create_body["customFields"] = plan["custom_fields"]

        status1, contact_id, detail1 = create_contact(create_body, dry_run=dry_run)
        if status1 == "failed":
            return "failed", None, detail1
        if dry_run:
            return "ok-would-create", None, f"{detail1}; note len={len(plan['note_body'])}"

        status3, detail3 = post_note_idempotent(contact_id, NOTE_TITLE, plan["note_body"], dry_run=dry_run)
        if status3 == "failed":
            return "failed", contact_id, detail3
        return "ok-created", contact_id, f"{detail1}; {detail3}"


def process_row(row_dict, notes_idx, contacts_idx, dry_run=False, fix_pass=False):
    try:
        plan = build_plan(row_dict, notes_idx, contacts_idx, fix_pass=fix_pass)
        return execute_plan(plan, dry_run=dry_run)
    except Exception as e:
        return "failed", None, f"exception: {e}"


def load_book():
    wb = openpyxl.load_workbook(BOOK_PATH, read_only=True, data_only=True)
    ws = wb["Book"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    all_rows = list(ws.iter_rows(min_row=2, values_only=True))
    notes_idx = build_notes_index(wb)
    contacts_idx = build_contacts_index(wb)
    return headers, all_rows, notes_idx, contacts_idx


def main():
    args = sys.argv[1:]
    fix = "--fix" in args
    args = [a for a in args if a != "--fix"]
    mode = args[0] if args else "dryrun10"
    plan_n = None
    if mode == "plan":
        plan_n = int(args[1]) if len(args) > 1 else 10

    headers, all_rows, notes_idx, contacts_idx = load_book()

    done = load_csv_log()

    candidates = []
    for i, row in enumerate(all_rows, start=2):
        rd = dict(zip(headers, row))
        if rd["Proposed GHL action"] in ("create", "update"):
            candidates.append((i, row))

    if mode == "plan":
        targets = candidates[:plan_n]
        print(f"PLAN (read-only, writes nothing): {len(targets)} of {len(candidates)} candidate rows")
        for row_num, row in targets:
            rd = dict(zip(headers, row))
            company = rd["Company"]
            plan = build_plan(rd, notes_idx, contacts_idx)
            action = "UPDATE existing" if plan["existing"] else "CREATE new"
            match = f" (matched by {plan['match_method']})" if plan["existing"] else ""
            fields = {k: v for k, v in {
                "firstName": plan["first"], "lastName": plan["last"], "email": plan["email"],
                "phone": plan["phone_norm"], "companyName": plan["company_name"],
                "address1": plan["addr1"], "city": plan["city"], "state": plan["state"],
                "postalCode": plan["postal"], "website": plan["website"],
            }.items() if v}
            note_dupe = ""
            if plan["existing"]:
                note_dupe = " [note already exists, would skip]" if has_note_with_title(plan["existing"]["id"], NOTE_TITLE) else " [note would be added]"
            print(f"row {row_num} | {company} | {action}{match} | fields={fields} | tags={plan['tags']} | note_len={len(plan['note_body'])}{note_dupe}")
        return

    if mode == "dryrun10":
        named = ["AAMCOR", "Cirque Lodge", "Ibex Plumbing", "5G Hearing"]
        picked = []
        for i, row in candidates:
            rd = dict(zip(headers, row))
            if any(n.lower() in (rd["Company"] or "").lower() for n in named):
                picked.append((i, row))
        others = [c for c in candidates if c not in picked]
        picked.extend(others[:10 - len(picked)])
        targets = picked[:10]
    elif fix:
        targets = [(i, row) for i, row in candidates
                   if dict(zip(headers, row))["Company"] in FIX_COMPANIES]
    else:
        targets = [c for c in candidates if dict(zip(headers, c[1]))["Company"] not in done]

    print(f"Processing {len(targets)} rows (mode={mode}, fix={fix}, total candidates={len(candidates)}, already done={len(done)})")
    for row_num, row in targets:
        rd = dict(zip(headers, row))
        company = rd["Company"]
        result_kind, contact_id, detail = process_row(rd, notes_idx, contacts_idx, fix_pass=fix)
        append_log(company, company, rd["Proposed GHL action"], contact_id, f"{result_kind}: {detail}")
        print(f"row {row_num} | {company} | {result_kind} | {contact_id} | {str(detail)[:160]}")


if __name__ == "__main__":
    main()
