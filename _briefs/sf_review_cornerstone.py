# Applies the Cornerstone CRM lookup (2026-09-28, read only, Dynamics accounts) to the last
# Needs your review rows of master book v5. Rules:
#   Cornerstone Active / CSR Assigned / Pending payroll / Not onboarded -> Cornerstone Client
#   Terminated -> Former Client ; Lost / Proposal Declined -> Lost
#   Not in either CRM: folder sits in a Lost folder -> Lost ; otherwise Former Client (unverified)
#   Active in both PEOs -> Cornerstone Client (Cornerstone effective date is newer) but kept for David to confirm
import openpyxl, re, shutil, json, sys
BOOK = sys.argv[1]
shutil.copy(BOOK, BOOK.replace('.xlsx', '-before-cs-lookup.xlsx'))
CS = {
 'ATL Fitness': ('Cornerstone Client', 'Cornerstone CRM: ATL Fitness LLC, Active, effective 2024-06-18'),
 'Centre for': ('Cornerstone Client', 'Cornerstone CRM: Active, effective 2026-03-03. Also active in G&A Salesforce'),
 'Centriforce Products': ('Cornerstone Client', 'Cornerstone CRM: Active, effective 2026-04-23. Also active in G&A Salesforce'),
 'CornerStone Cunstruction': ('Lost', 'Cornerstone CRM: Cornerstone Concrete N Construction LLC, Proposal Declined'),
 'Fat Karting': ('Cornerstone Client', 'Cornerstone CRM: FAT Karting League LLC (plus CA and TX entities), Active, effective 2025-10-24'),
 'Good Bones': ('Cornerstone Client', 'Cornerstone CRM: Active, effective 2025-12-26'),
 'HBC Homes': ('Cornerstone Client', 'Cornerstone CRM: Active, effective 2026-08-05. Also active in G&A Salesforce'),
 'KW Powder': ('Cornerstone Client', 'Cornerstone CRM: KW Powder & Fab LLC, CSR Assigned, pending initial payroll, effective 2026-04-07'),
 'Mery Floor': ('Cornerstone Client', 'Cornerstone CRM: Mery Floor Coverings Inc, Active, not onboarded yet, effective 2025-12-30'),
 'Mortensen Orthodontics': ('Cornerstone Client', 'Cornerstone CRM: Active, effective 2025-07-16'),
 'North Shore': ('Former Client', 'Cornerstone CRM: North Shore Lodge LLC, Terminated 2025-06-23'),
 'Realeyes': ('Cornerstone Client', 'Cornerstone CRM: Realeyes OU LLC, Active, pending initial payroll, effective 2026-06-05'),
 'Rhino Insurance': ('Cornerstone Client', 'Cornerstone CRM: CSR Assigned, Active, effective 2025-08-27. Also active in G&A Salesforce'),
 'Rider Telecom': ('Former Client', 'Cornerstone CRM: Rider Telecom Inc, Terminated 2024-12-20'),
 'Shufti Pro': ('Cornerstone Client', 'Cornerstone CRM: Shufti Pro Delaware LLC, Active, effective 2026-04-13'),
 'Valley Construction': ('Lost', 'Cornerstone CRM: Valley Construction of Nevada LLC, Lost'),
 'Velvet Ink': ('Cornerstone Client', 'Cornerstone CRM: Active, effective 2025-12-26'),
 'Wilson Custom': ('Cornerstone Client', 'Cornerstone CRM: Wilson Custom Builders LLC, Active, effective 2025-06-28'),
}
BOTH = {'Centre for', 'Centriforce Products', 'HBC Homes', 'Rhino Insurance'}
stop = {'llc','inc','inc.','llc.','co','corp','the','and','&','dba','of','company','group','services','pllc','pc','ltd','l.l.c.','-'}
def key(n):
    w = [x for x in re.split(r'[\s,.()]+', n) if x and x.lower() not in stop]
    k = ' '.join(w[:2]).replace("'", '')
    for c in CS:
        if k.lower() == c.lower() or (c == 'Realeyes' and k.startswith('Realeyes')): return c
    return None
wb = openpyxl.load_workbook(BOOK)
rv, bk = wb['Needs your review'], wb['Book']
ix = {c.value: c.column for c in bk[1]}
rvrows = [[c.value for c in row] for row in rv.iter_rows(min_row=2)]
rvhdr = [c.value for c in rv[1]]
byco = {}
for r in range(2, bk.max_row + 1): byco.setdefault(bk.cell(r, ix['Company']).value, []).append(r)
resolved, keep, changed = [], [], []
for row in rvrows:
    comp, folders = row[0], (row[6] or '')
    k = key(comp)
    if k:
        new, why = CS[k]
        active_with = 'Cornerstone' if new == 'Cornerstone Client' else None
    elif re.search(r'lost', folders, re.I):
        new, why, active_with = 'Lost', 'Not in Cornerstone CRM or G&A Salesforce; folder sits in a Lost folder', None
    else:
        new, why, active_with = 'Former Client', 'Unverified: client folder exists but no account in Cornerstone CRM or G&A Salesforce (checked 2026-09-28)', None
    for r in byco.get(comp, []):
        old = bk.cell(r, ix['Status']).value
        bk.cell(r, ix['Status']).value = new
        bk.cell(r, ix['Final status (all systems)']).value = new
        bk.cell(r, ix['Why']).value = why
        if active_with: bk.cell(r, ix['Active with']).value = active_with
        if k in BOTH:
            bk.cell(r, ix['Needs your review']).value = 'Active in both PEOs. Set to Cornerstone (newer effective date). Confirm.'
        else:
            bk.cell(r, ix['Needs your review']).value = None
        changed.append((comp, old, new, bk.cell(r, ix['GHL Contact ID']).value))
    if k in BOTH: keep.append(row)
    else: resolved.append((comp, new, why))
for rr in range(rv.max_row, 1, -1): rv.delete_rows(rr)
for row in keep:
    row[5] = 'Active in both PEOs. I set Cornerstone because its effective date is newer. Right?'
    rv.append(row)
res = wb['Resolved by lookup']
for x in sorted(resolved): res.append(list(x))
wb.save(BOOK)
json.dump(changed, open('/tmp/cs_lookup_changed.json', 'w'))
import collections
print('resolved', len(resolved), 'kept', len(keep))
print(collections.Counter(c[2] for c in changed))
print(collections.Counter(bk.cell(r, ix['Status']).value for r in range(2, bk.max_row + 1)))
print('with GHL id', sum(1 for c in changed if c[3]))
